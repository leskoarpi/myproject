"""Excel export for the monthly worksheets.

Produces a real .xlsx (openpyxl), not a CSV with a spreadsheet extension, so
the sheet opens in Excel, LibreOffice or Google Sheets with the header frozen,
the weekend columns shaded and the day columns narrow enough to fit a month on
one screen.
"""

import io
import logging

logging.getLogger(__name__)

HEADER_FILL = "1F4E79"
WEEKEND_FILL = "E8EEF4"
BORDER_COLOR = "B8C4D0"


def grid_to_xlsx(grid):
    """Render a :class:`~apps.reports.monthly.MonthlyGrid` as .xlsx bytes.

    Returns ``None`` when openpyxl is not installed, so the caller can fall
    back to the printable HTML view rather than erroring.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:  # pragma: no cover - optional dependency
        return None

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"{grid.title} {grid.year}-{grid.month:02d}"[:31]

    thin = Side(style="thin", color=BORDER_COLOR)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    weekend_fill = PatternFill("solid", fgColor=WEEKEND_FILL)
    centered = Alignment(horizontal="center", vertical="center")

    # --- title row ---
    label_count = len(grid.row_headers)
    total_columns = label_count + len(grid.days) + len(grid.summary_headers)
    sheet.cell(row=1, column=1, value=f"{grid.title} — {grid.period_label}").font = Font(
        bold=True, size=13
    )
    sheet.merge_cells(
        start_row=1, start_column=1, end_row=1, end_column=max(total_columns, 2)
    )

    # --- header row ---
    header_row = 3
    columns = list(grid.row_headers) + [str(d) for d in grid.days] + list(
        grid.summary_headers
    )
    for index, title in enumerate(columns, start=1):
        cell = sheet.cell(row=header_row, column=index, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = centered
        cell.border = border

    weekend = grid.weekend_days

    # --- data rows ---
    for offset, row in enumerate(grid.rows, start=header_row + 1):
        column = 1
        for label in row["labels"]:
            cell = sheet.cell(row=offset, column=column, value=label)
            cell.border = border
            column += 1
        for day in grid.days:
            cell = sheet.cell(row=offset, column=column, value=row["cells"].get(day, ""))
            cell.alignment = centered
            cell.border = border
            if day in weekend:
                cell.fill = weekend_fill
            column += 1
        for value in row.get("summary", []):
            cell = sheet.cell(row=offset, column=column, value=value)
            cell.alignment = centered
            cell.border = border
            cell.font = Font(bold=True, size=10)
            column += 1

    # --- legend ---
    if grid.legend:
        legend_row = header_row + len(grid.rows) + 2
        sheet.cell(row=legend_row, column=1, value="Jelmagyarázat:").font = Font(bold=True)
        for index, (code, meaning) in enumerate(grid.legend, start=1):
            sheet.cell(row=legend_row + index, column=1, value=code)
            sheet.cell(row=legend_row + index, column=2, value=meaning)

    # --- sizing: wide labels, narrow days, everything visible while scrolling ---
    widths = [10, 26, 10][:label_count] or [12] * label_count
    while len(widths) < label_count:
        widths.append(12)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for index in range(label_count + 1, label_count + len(grid.days) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 4.2
    for index in range(
        label_count + len(grid.days) + 1, total_columns + 1
    ):
        sheet.column_dimensions[get_column_letter(index)].width = 11

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=label_count + 1)

    # --- print setup: A4 landscape, header repeated on every page ---
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = f"{header_row}:{header_row}"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
