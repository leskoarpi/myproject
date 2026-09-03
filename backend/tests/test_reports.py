"""Monthly worksheet reports: room tidiness and evening presence.

Both are grids - one row per room or student, one column per day of the month -
rendered as a printable page and as a real .xlsx worksheet.
"""

import datetime as dt
import io

import pytest
from django.urls import reverse

from apps.accounts.capabilities import Role
from apps.dormcalendar.services import get_or_create_day
from apps.inspections.services import (
    open_evening_session,
    open_room_check_session,
    save_evening_result,
    save_room_check,
)
from apps.reports.monthly import build_grid, evening_presence_grid, room_tidiness_grid
from apps.reports.xlsx import grid_to_xlsx

from . import factories as f

pytestmark = pytest.mark.django_db

YEAR, MONTH = 2026, 9


@pytest.fixture
def setup():
    f.seed_status_types()
    f.seed_modules()
    year = f.make_school_year()
    room = f.make_room(number="101", floor=1, capacity=4)
    student = f.make_student(name="Teszt Diák")
    f.assign(student, room, school_year=year)
    return {
        "admin": f.make_user(Role.ADMIN),
        "room": room,
        "student": student,
        "year": year,
    }


def _rate(setup, day, rating):
    session = open_room_check_session(date=day, floor=1, actor=setup["admin"])
    save_room_check(room_check=session.checks.first(), actor=setup["admin"], rating=rating)


def _evening(setup, day, status_code):
    calendar_day = get_or_create_day(day)
    calendar_day.evening_check_required = True
    calendar_day.save()
    session = open_evening_session(calendar_day=calendar_day, floor=1, actor=setup["admin"])
    save_evening_result(
        session=session, student=setup["student"], status=status_code, actor=setup["admin"]
    )


# --------------------------------------------------------------------------
# Grid shape
# --------------------------------------------------------------------------


def test_grid_has_one_column_per_day_of_the_month(setup):
    grid = room_tidiness_grid(YEAR, MONTH)
    assert grid.days == list(range(1, 31))  # September has 30 days


def test_february_length_follows_the_calendar(setup):
    assert len(room_tidiness_grid(2024, 2).days) == 29  # leap year
    assert len(room_tidiness_grid(2026, 2).days) == 28


def test_weekend_columns_are_marked(setup):
    grid = room_tidiness_grid(YEAR, MONTH)
    # 2026-09-05 is a Saturday, 2026-09-07 a Monday.
    assert 5 in grid.weekend_days
    assert 7 not in grid.weekend_days


def test_tidiness_grid_places_ratings_on_the_right_days(setup):
    _rate(setup, dt.date(YEAR, MONTH, 2), 5)
    _rate(setup, dt.date(YEAR, MONTH, 9), 3)

    grid = room_tidiness_grid(YEAR, MONTH)
    row = next(r for r in grid.rows if r["labels"][0] == "101")

    assert row["cells"][2] == 5
    assert row["cells"][9] == 3
    assert 3 not in row["cells"]
    assert row["summary"] == [4.0, 2]  # average, number of checks


def test_tidiness_grid_ignores_other_months(setup):
    _rate(setup, dt.date(YEAR, MONTH, 2), 5)
    _rate(setup, dt.date(YEAR, 8, 20), 1)

    row = next(r for r in room_tidiness_grid(YEAR, MONTH).rows if r["labels"][0] == "101")
    assert list(row["cells"]) == [2]


def test_evening_grid_uses_short_status_codes(setup):
    _evening(setup, dt.date(YEAR, MONTH, 1), "inside")
    _evening(setup, dt.date(YEAR, MONTH, 2), "home")

    grid = evening_presence_grid(YEAR, MONTH)
    row = next(r for r in grid.rows if "Teszt Diák" in r["labels"])

    assert row["cells"][1] == "B"
    assert row["cells"][2] == "H"
    assert row["summary"] == [1, 1, 2]  # inside, other, checked
    assert row["labels"][0] == "101"


def test_empty_month_still_renders_a_grid(setup):
    grid = evening_presence_grid(YEAR, MONTH)
    assert grid.rows == []
    assert len(grid.days) == 30


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------


def test_xlsx_is_a_real_workbook(setup):
    _rate(setup, dt.date(YEAR, MONTH, 2), 4)
    payload = grid_to_xlsx(room_tidiness_grid(YEAR, MONTH))

    assert payload is not None
    assert payload[:2] == b"PK"  # xlsx is a zip container

    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(payload)).active
    assert "Szobarend" in sheet.title
    # Header row: labels, then day numbers.
    assert sheet.cell(row=3, column=1).value == "Szoba"
    assert sheet.cell(row=3, column=3).value == "1"
    assert sheet.freeze_panes == "C4"
    assert sheet.page_setup.orientation == "landscape"


def test_xlsx_carries_the_ratings(setup):
    _rate(setup, dt.date(YEAR, MONTH, 2), 4)
    from openpyxl import load_workbook

    sheet = load_workbook(
        io.BytesIO(grid_to_xlsx(room_tidiness_grid(YEAR, MONTH)))
    ).active
    # Row 4 is the first data row; day 2 sits in column 2 (labels) + 2.
    assert sheet.cell(row=4, column=1).value == "101"
    assert sheet.cell(row=4, column=4).value == 4


# --------------------------------------------------------------------------
# Views
# --------------------------------------------------------------------------


def test_printable_sheet_renders(client, setup):
    _rate(setup, dt.date(YEAR, MONTH, 2), 4)
    client.force_login(setup["admin"])

    response = client.get(
        reverse("reports:monthly", kwargs={"name": "room_tidiness"}),
        {"year": YEAR, "month": MONTH},
    )
    body = response.content.decode()

    assert response.status_code == 200
    assert "Szobarend" in body
    assert "101" in body


def test_xlsx_download_has_the_right_content_type(client, setup):
    client.force_login(setup["admin"])
    response = client.get(
        reverse("reports:monthly_xlsx", kwargs={"name": "evening_presence"}),
        {"year": YEAR, "month": MONTH},
    )

    assert response.status_code == 200
    assert response["Content-Type"].endswith("spreadsheetml.sheet")
    assert ".xlsx" in response["Content-Disposition"]


def test_export_is_audited(client, setup):
    from apps.audit.models import AuditAction, AuditLog

    client.force_login(setup["admin"])
    client.get(
        reverse("reports:monthly_xlsx", kwargs={"name": "room_tidiness"}),
        {"year": YEAR, "month": MONTH},
    )
    assert AuditLog.objects.filter(action=AuditAction.EXPORT).exists()


def test_unknown_sheet_is_404(client, setup):
    client.force_login(setup["admin"])
    assert (
        client.get(reverse("reports:monthly", kwargs={"name": "nonsense"})).status_code == 404
    )


def test_a_teacher_cannot_reach_the_reports(client, setup):
    teacher = f.make_teacher()
    client.force_login(teacher.user)
    assert client.get(reverse("reports:index")).status_code == 403


def test_build_grid_rejects_unknown_names(setup):
    with pytest.raises(KeyError):
        build_grid("nope", YEAR, MONTH)
